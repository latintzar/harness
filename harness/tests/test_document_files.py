"""Real Office/PDF rendering, text-preview fidelity and bounded conversion."""
from io import BytesIO
import hashlib
import json
from pathlib import Path
import sqlite3
import subprocess
import sys
from zipfile import ZipFile

import pytest

sys.path.insert(0, str(Path(__file__).parents[1]))
from chat_files import ChatFiles
import document_rendering

SCOPE = {'owner_principal_id': 'alice', 'chat_id': 'chat-a', 'dataset': 'cafe'}
SOURCE = '# Service report\n\nCafé София: **44 €**.\n\n## Follow-up\n\n- Check the counter.\n- Review the report.\n\n| Area | Action |\n| --- | --- |\n| Counter | Review |\n'


def create(tmp_path, name, source):
    store = ChatFiles(tmp_path)
    record = store.create({'name': name, 'content': source}, **SCOPE)
    metadata, binary = store.download(record['id'], **SCOPE)
    assert metadata == record and len(binary) == record['size_bytes']
    assert hashlib.sha256(binary).hexdigest() == record['sha256']
    assert store.content(record['id'], **SCOPE)['content'] == source
    assert record['source_sha256'] == hashlib.sha256(source.encode()).hexdigest()
    return record, binary


def test_docx_has_real_a4_page_and_preserves_headings_tables_and_unicode(tmp_path):
    from docx import Document
    record, binary = create(tmp_path, 'report.docx', SOURCE)
    assert record['display_type'] == 'document' and record['preview_format'] == 'markdown'
    document = Document(BytesIO(binary))
    section = document.sections[0]
    # Word stores page dimensions in integer twips; A4 rounds within 0.02 mm.
    assert section.page_width.mm == pytest.approx(210, abs=.02)
    assert section.page_height.mm == pytest.approx(297, abs=.02)
    assert document.styles['Normal'].font.name == 'Arial'
    assert [(p.style.name, p.text) for p in document.paragraphs if p.style.name.startswith('Heading')] == [
        ('Heading 1', 'Service report'), ('Heading 2', 'Follow-up')]
    assert 'Café София: 44 €.' in [p.text for p in document.paragraphs]
    assert any(run.bold and run.text == '44 €' for p in document.paragraphs for run in p.runs)
    assert document.tables[0].cell(1, 0).text == 'Counter'
    assert ChatFiles(tmp_path).download(record['id'], **SCOPE)[1] == binary


def test_pdf_is_a_real_unicode_document_with_embedded_font(tmp_path):
    from pypdf import PdfReader
    record, binary = create(tmp_path, 'report.pdf', SOURCE)
    assert record['display_type'] == 'pdf' and binary.startswith(b'%PDF-')
    document = PdfReader(BytesIO(binary))
    text = '\n'.join(page.extract_text() for page in document.pages)
    for expected in ('Service report', 'Café София:', '44 €', 'Counter', 'Review'):
        assert expected in text
    assert float(document.pages[0].mediabox.width) == pytest.approx(595.276, abs=.01)
    assert float(document.pages[0].mediabox.height) == pytest.approx(841.89, abs=.01)
    assert b'/FontFile2' in binary


def test_xlsx_preserves_tabular_content_without_executable_formulas(tmp_path):
    from openpyxl import load_workbook
    source = 'Name,Value\nCafé,44\nOrder ID,00042\nFormula,=HYPERLINK("https://example.com")\nBudget,-9.5\n'
    record, binary = create(tmp_path, 'report.xlsx', source)
    assert record['display_type'] == 'spreadsheet' and record['preview_format'] == 'table'
    workbook = load_workbook(BytesIO(binary), data_only=False)
    sheet = workbook.active
    assert sheet['A2'].value == 'Café' and sheet['B2'].value == 44
    assert sheet['B3'].value == '00042' and sheet['B3'].data_type == 's'
    assert sheet['B4'].value == '=HYPERLINK("https://example.com")' and sheet['B4'].data_type == 's'
    assert sheet['B5'].value == -9.5 and sheet['B5'].data_type == 'n'
    assert sheet.freeze_panes == 'A2' and sheet.auto_filter.ref == 'A1:B5'
    with ZipFile(BytesIO(binary)) as archive:
        assert b'<f>' not in archive.read('xl/worksheets/sheet1.xml')
        assert not any('externalLink' in name or name.endswith('vbaProject.bin') for name in archive.namelist())


def test_document_references_remain_inert_text(tmp_path):
    from docx import Document
    source = '# References\n\n<img src="http://127.0.0.1:1/private">\n\n![Camera](file:///etc/passwd)\n\n![Remote image](https://example.com/image.png)\n\n[Website](https://example.com)'
    _, binary = create(tmp_path, 'references.docx', source)
    document = Document(BytesIO(binary))
    text = '\n'.join(p.text for p in document.paragraphs)
    assert '<img src="http://127.0.0.1:1/private">' in text
    assert '![Camera](file:///etc/passwd)' in text and 'Website (https://example.com)' in text
    assert '[Image: Remote image]' in text
    with ZipFile(BytesIO(binary)) as archive:
        assert not any(name.startswith('word/media/') for name in archive.namelist())
        assert b'TargetMode="External"' not in archive.read('word/_rels/document.xml.rels')


def test_document_bounds_keep_failed_outputs_out_of_the_catalogue(tmp_path):
    store = ChatFiles(tmp_path)
    for name, source, message in [
        ('huge.docx', 'x' * (document_rendering.MAX_SOURCE_BYTES + 1), '128 KiB'),
        ('large.xlsx', 'Name\n' * (document_rendering.MAX_ROWS + 1), '2,000 rows'),
        ('control.docx', 'word\x01text', 'readable text'),
        ('emoji.pdf', '# Report\n🫠', 'Use DOCX'),
    ]:
        with pytest.raises(ValueError, match=message):
            store.create({'name': name, 'content': source}, **SCOPE)
    assert store.list(owner_principal_id='alice') == []


def test_conversion_timeout_releases_slot_and_never_inherits_credentials(tmp_path, monkeypatch):
    calls = []
    def run(args, **kwargs):
        calls.append((args, kwargs))
        assert 'OPENAI_API_KEY' not in kwargs['env'] and kwargs['timeout'] == 15
        assert kwargs['close_fds'] and kwargs['env']['OPENBLAS_NUM_THREADS'] == '1'
        raise subprocess.TimeoutExpired(args, 15)
    monkeypatch.setenv('OPENAI_API_KEY', 'fixture-private-key')
    with monkeypatch.context() as scoped:
        scoped.setattr(document_rendering.subprocess, 'run', run)
        with pytest.raises(ValueError, match='smaller sections'):
            create(tmp_path, 'report.docx', SOURCE)
    assert len(calls) == 1
    assert create(tmp_path, 'report.docx', SOURCE)[1].startswith(b'PK')


def test_existing_text_database_migrates_without_changing_download_bytes(tmp_path):
    path = tmp_path / 'chat-files.sqlite'
    identity = 'file-' + 'a' * 32
    metadata = {'id': identity, 'kind': 'document', 'chat_id': 'chat-a', 'dataset': 'cafe', 'name': 'old.txt'}
    with sqlite3.connect(path) as db:
        db.execute('CREATE TABLE files (id TEXT PRIMARY KEY, owner TEXT, chat_id TEXT, dataset TEXT, metadata TEXT, content BLOB)')
        db.execute('INSERT INTO files VALUES(?,?,?,?,?,?)', (identity, 'alice', 'chat-a', 'cafe', json.dumps(metadata), b'Original bytes'))
    store = ChatFiles(tmp_path)
    assert store.download(identity, **SCOPE) == (metadata, b'Original bytes')
    with pytest.raises(ValueError):
        store.download(identity, owner_principal_id='bob')
