"""Trusted generated documents; all Office/PDF parsing happens in Docker."""
import hashlib
import json
import os
from pathlib import Path
import subprocess
import uuid
import zipfile

import pytest

IMAGE=os.environ.get('DOCUMENT_PREVIEW_TEST_IMAGE')
pytestmark=pytest.mark.skipif(not IMAGE,reason='Set the locally built isolated fixture image.')
FILE_ID='file-'+'a'*32

def run_worker(tmp_path,original,format,*,timeout_fixture=False):
    output=tmp_path/'output';output.mkdir(mode=0o777);output.chmod(0o777)
    original.chmod(0o444);name='rocoto-office-test-'+uuid.uuid4().hex[:12]
    command=['docker','run','--name',name,'--network','none','--read-only','--cap-drop','ALL',
        '--security-opt','no-new-privileges','--user','10001:10001','--cpus','2',
        '--memory','1g','--memory-swap','1g','--pids-limit','64','--init',
        '--tmpfs','/tmp:rw,noexec,nosuid,nodev,size=64m',
        '--mount',f'type=bind,src={original},dst=/input/original,readonly',
        '--mount',f'type=bind,src={output},dst=/output']
    if timeout_fixture:
        code=f"from document_preview_worker import prepare,Limits; prepare({FILE_ID!r},{format!r},limits=Limits(job_seconds=.001,command_seconds=.001))"
        command+=['--entrypoint','python',IMAGE,'-c',code]
    else:command+=[IMAGE,'--file-id',FILE_ID,'--format',format]
    try:return subprocess.run(command,capture_output=True,text=True,timeout=150),output
    finally:subprocess.run(['docker','rm','--force',name],check=True,stdout=subprocess.DEVNULL)

def verify(output,original,pages):
    root=output/'preview';m=json.loads((root/'manifest.json').read_text())
    assert m['original_sha256']==hashlib.sha256(original.read_bytes()).hexdigest()
    assert m['original_size_bytes']==original.stat().st_size and m['page_count']==pages
    assert len(m['pages'])==pages and m['file_id']==FILE_ID
    for row in [*m['pages'],m['text']]:
        raw=(root/row['path']).read_bytes();assert hashlib.sha256(raw).hexdigest()==row['sha256']
        assert len(raw)==row['size_bytes']
    assert set(x.name for x in output.iterdir())=={'preview'}
    return m,(root/'text.txt').read_text()

def test_docx_actual_two_pages_and_literal_text(tmp_path):
    from docx import Document
    d=Document();d.add_heading('First document page',0);d.add_paragraph('Exact original content: Café — 123.')
    d.add_table(rows=2,cols=2).cell(0,0).text='Visible table';d.add_page_break();d.add_heading('Second document page',0)
    d.add_paragraph('This stays on page two.');original=tmp_path/'report.docx';d.save(original)
    result,out=run_worker(tmp_path,original,'docx');assert result.returncode==0,result.stdout+result.stderr
    manifest,text=verify(out,original,2)
    assert 'First document page' in text and 'Second document page' in text and 'Visible table' in text
    assert not manifest['text']['truncated']

def test_pdf_actual_two_pages(tmp_path):
    from reportlab.pdfgen.canvas import Canvas
    original=tmp_path/'report.pdf';c=Canvas(str(original));c.drawString(72,700,'PDF first page');c.showPage();c.drawString(72,700,'PDF second page');c.save()
    result,out=run_worker(tmp_path,original,'pdf');assert result.returncode==0,result.stdout+result.stderr
    _,text=verify(out,original,2);assert 'PDF first page' in text and 'PDF second page' in text

def test_xlsx_actual_printed_sheet(tmp_path):
    from openpyxl import Workbook
    book=Workbook();sheet=book.active;sheet.title='Sales';sheet.append(['Product','Units']);sheet.append(['Fixture',12]);sheet.print_area='A1:B2'
    original=tmp_path/'report.xlsx';book.save(original)
    result,out=run_worker(tmp_path,original,'xlsx');assert result.returncode==0,result.stdout+result.stderr
    _,text=verify(out,original,1);assert 'Product' in text and 'Fixture' in text and '12' in text

def test_invalid_pdf_fails_without_partial_preview(tmp_path):
    original=tmp_path/'report.pdf';original.write_bytes(b'%PDF-invalid')
    result,out=run_worker(tmp_path,original,'pdf');assert result.returncode==2
    assert json.loads(result.stdout)['status']=='failed' and not list(out.iterdir())

def test_external_office_dependency_rejected_inside_worker(tmp_path):
    from docx import Document
    original=tmp_path/'report.docx';d=Document();d.add_paragraph('External fixture');d.save(original)
    with zipfile.ZipFile(original,'a') as z:z.writestr('word/_rels/external.rels','<Relationships><Relationship TargetMode="External" Type="image" Target="https://example.com/image.png"/></Relationships>')
    result,out=run_worker(tmp_path,original,'docx');assert result.returncode==2
    assert json.loads(result.stdout)['code']=='document_external_content' and not list(out.iterdir())

def test_page_limit_has_truthful_failure(tmp_path):
    from reportlab.pdfgen.canvas import Canvas
    original=tmp_path/'report.pdf';c=Canvas(str(original))
    for page in range(41):c.drawString(72,700,str(page));c.showPage()
    c.save();result,out=run_worker(tmp_path,original,'pdf');assert result.returncode==2
    assert json.loads(result.stdout)['code']=='document_page_limit' and not list(out.iterdir())

def test_macro_archive_rejected_inside_worker(tmp_path):
    from docx import Document
    original=tmp_path/'report.docx';d=Document();d.add_paragraph('Macro fixture');d.save(original)
    with zipfile.ZipFile(original,'a') as archive:archive.writestr('word/vbaProject.bin',b'fixture-macro')
    result,out=run_worker(tmp_path,original,'docx');assert result.returncode==2
    assert json.loads(result.stdout)['code']=='document_active_content' and not list(out.iterdir())

def test_timeout_reaps_conversion_and_discards_partial_output(tmp_path):
    from reportlab.pdfgen.canvas import Canvas
    original=tmp_path/'report.pdf';c=Canvas(str(original));c.drawString(72,700,'Timeout fixture');c.save()
    result,out=run_worker(tmp_path,original,'pdf',timeout_fixture=True)
    assert result.returncode!=0 and 'preview_timeout' in result.stderr and not list(out.iterdir())

def test_encrypted_pdf_has_useful_failure_and_original_remains_exact(tmp_path):
    from reportlab.pdfgen.canvas import Canvas
    original=tmp_path/'report.pdf';c=Canvas(str(original),encrypt='synthetic-password');c.drawString(72,700,'Private');c.save()
    before=original.read_bytes();result,out=run_worker(tmp_path,original,'pdf')
    assert result.returncode==2 and original.read_bytes()==before and not list(out.iterdir())

def test_create_file_actual_docx_download_renders_and_source_stays_exact(tmp_path):
    from chat_files import ChatFiles
    files=ChatFiles(tmp_path/'state');source='# Generated report\n\nExact generated content **stays available**.'
    record=files.create({'name':'Generated.docx','title':'Generated report','content':source},owner_principal_id='alice',chat_id='root',dataset='company')
    before,raw=files.download(record['id'],owner_principal_id='alice')
    original=tmp_path/'generated.docx';original.write_bytes(raw)
    result,out=run_worker(tmp_path,original,'docx');assert result.returncode==0,result.stdout+result.stderr
    _,text=verify(out,original,1)
    assert 'Exact generated content stays available.' in text
    assert files.download(record['id'],owner_principal_id='alice')==(before,raw)
    assert files.content(record['id'],owner_principal_id='alice')['content']==source
